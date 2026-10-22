"""
Export the Markdown cheat sheet to DOCX and PDF.

This script intentionally supports a small subset of Markdown:
- #, ##, ### headings
- bullet lines that start with "- "
- numbered lines like "1. "
- plain paragraphs

It is enough for the included analyst cheat sheet and can be reused.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "docs" / "Analyst_Cheat_Sheet.md"
OUT_DOCX = ROOT / "docs" / "Analyst_Cheat_Sheet.docx"
OUT_PDF = ROOT / "docs" / "Analyst_Cheat_Sheet.pdf"


def read_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


def add_to_docx(lines: list[str], out_path: Path) -> None:
    doc = Document()
    for raw in lines:
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        if re.match(r"^\d+\.\s", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            doc.add_paragraph(text, style="List Number")
            continue

        doc.add_paragraph(line)

    doc.save(out_path)


def wrap_text_for_pdf(text: str, max_chars: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    cur = words[0]
    for w in words[1:]:
        if len(cur) + 1 + len(w) <= max_chars:
            cur = f"{cur} {w}"
        else:
            lines.append(cur)
            cur = w
    lines.append(cur)
    return lines


def add_to_pdf(lines: list[str], out_path: Path) -> None:
    c = canvas.Canvas(str(out_path), pagesize=LETTER)
    width, height = LETTER

    left = 54
    y = height - 54
    base_size = 11
    line_step = 15

    def new_page_if_needed(current_y: float) -> float:
        if current_y < 54:
            c.showPage()
            return height - 54
        return current_y

    for raw in lines:
        line = raw.rstrip()
        if not line:
            y -= line_step
            y = new_page_if_needed(y)
            continue

        # Basic style handling from markdown markers.
        font = "Helvetica"
        size = base_size
        prefix = ""
        content = line

        if line.startswith("# "):
            font = "Helvetica-Bold"
            size = 18
            content = line[2:].strip()
        elif line.startswith("## "):
            font = "Helvetica-Bold"
            size = 14
            content = line[3:].strip()
        elif line.startswith("### "):
            font = "Helvetica-Bold"
            size = 12
            content = line[4:].strip()
        elif line.startswith("- "):
            prefix = "- "
            content = line[2:].strip()
        elif re.match(r"^\d+\.\s", line):
            prefix = re.match(r"^\d+\.\s", line).group(0)
            content = re.sub(r"^\d+\.\s+", "", line)

        max_chars = 95 if size <= 12 else 70
        wrapped = wrap_text_for_pdf(prefix + content, max_chars=max_chars)

        c.setFont(font, size)
        for wline in wrapped:
            y = new_page_if_needed(y)
            c.drawString(left, y, wline)
            y -= (line_step + (2 if size > 12 else 0))

    c.save()


def main() -> None:
    lines = read_lines(SRC_MD)
    add_to_docx(lines, OUT_DOCX)
    add_to_pdf(lines, OUT_PDF)
    print(f"Created: {OUT_DOCX}")
    print(f"Created: {OUT_PDF}")


if __name__ == "__main__":
    main()
