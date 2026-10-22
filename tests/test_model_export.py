"""
Model exports (`V6_PLAN.md` Phase D).

This is the half that makes modelling here worth doing rather than exporting to
Stata, so what is under test is not "does a file appear" — it is whether the
document that reaches a chapter still carries what qualifies the numbers.

Three readers, three shapes: Word for the publishable table, Excel for working,
CSV plus a sidecar for anyone who wants R. Each is checked for the thing it
would be easiest to leave out.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from analysis.design_matrix import DesignMatrix
from analysis.models import fit
from core.model_spec import ModelRequest, Shelf
from exports.model_export import (
    SHEET_CAVEATS, SHEET_ESTIMATES, SHEET_SPEC, build_model_docx,
    build_model_workbook, design_matrix_sidecar, export_design_matrix)


def matrix(n=300, seed=5, machine=False):
    rng = np.random.default_rng(seed)
    cost = rng.integers(0, 2, n)
    first_gen = rng.integers(0, 2, n)
    gpa = rng.normal(3.0, 0.5, n)
    outcome = 1.0 * cost - 0.4 * first_gen + rng.normal(scale=0.5, size=n)
    frame = pd.DataFrame({
        "case_id": [str(i) for i in range(n)],
        "code::Cost": cost, "code::Advising": rng.integers(0, 2, n),
        "desc::first_gen": first_gen, "desc::gpa": gpa,
        "wellbeing": outcome})
    rater = "transformer" if machine else "human_primary"
    return DesignMatrix(frame=frame, unit="source_id", encoding="binary",
                        rater_id=rater, raters_pooled=[rater],
                        reference_levels={"program": "Arts"},
                        dropped={"code::Rare": "varies on only 2 case(s); 6 needed"})


def request(method="lasso", treatment=()):
    """A spec. A named treatment is removed from Predictors, because a field on
    both shelves makes its effect uninterpretable and `problems()` blocks it —
    which the first version of this helper did not do, and the export then
    happily wrote a document for the refused model."""
    predictors = tuple(
        name for name in ("code::Cost", "code::Advising", "desc::first_gen",
                          "desc::gpa")
        if name not in treatment)
    return ModelRequest(
        method=method, outcome=Shelf(("wellbeing",)),
        treatment=Shelf(tuple(treatment)),
        predictors=Shelf(predictors))


def docx_text(path):
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.fixture()
def fitted():
    built = matrix()
    spec = request()
    return fit(built, spec), spec, built


class TestTheWordTableCarriesItsCaveats:
    """A coefficient table looks authoritative whatever produced it, and the
    person reading it in a chapter will not have the program open."""

    def test_the_estimates_reach_the_document(self, fitted, tmp_path):
        result, spec, built = fitted

        path = build_model_docx(result, spec, tmp_path / "m.docx", matrix=built)
        text = docx_text(path)

        assert "code::Cost" in text

    def test_every_caveat_reaches_the_document(self, fitted, tmp_path):
        result, spec, built = fitted

        text = docx_text(build_model_docx(result, spec, tmp_path / "m.docx",
                                          matrix=built))

        for note in result.notes:
            assert note[:60] in text, f"missing caveat: {note[:60]}"

    def test_the_seed_is_printed(self, fitted, tmp_path):
        result, spec, built = fitted

        text = docx_text(build_model_docx(result, spec, tmp_path / "m.docx",
                                          matrix=built))

        assert str(result.seed) in text

    def test_a_reviewer_could_reconstruct_the_fit(self, fitted, tmp_path):
        """The list that makes this worth exporting at all."""
        result, spec, built = fitted

        text = docx_text(build_model_docx(result, spec, tmp_path / "m.docx",
                                          matrix=built))

        assert "Outcome: wellbeing" in text
        assert "Cases modelled" in text
        assert "Random seed" in text
        assert "reference levels" in text.lower()
        assert "code::Rare" in text, "the dropped column is not recorded"

    def test_machine_coded_variables_are_declared(self, tmp_path):
        """A coefficient on a machine-coded variable has to say so in the
        document, not only on the screen where it was fitted."""
        built = matrix(machine=True)
        spec = request()

        text = docx_text(build_model_docx(fit(built, spec), spec,
                                          tmp_path / "m.docx", matrix=built))

        assert "estimates, not measurements" in text

    def test_an_importance_table_says_it_has_no_direction(self, tmp_path):
        built = matrix()
        spec = request("random_forest")

        text = docx_text(build_model_docx(fit(built, spec), spec,
                                          tmp_path / "m.docx", matrix=built))

        assert "no direction" in text
        assert "Variable importance" in text

    def test_a_tree_exports_its_rules(self, tmp_path):
        built = matrix()
        spec = request("decision_tree")

        text = docx_text(build_model_docx(fit(built, spec), spec,
                                          tmp_path / "m.docx", matrix=built))

        assert "Rules" in text
        assert "code::Cost" in text

    def test_a_refused_model_produces_no_document(self, tmp_path):
        """A file gets circulated; a refusal on screen does not. The refusal is
        the finding and belongs there, not in a Word document with a title."""
        from analysis.models import ModelResult

        from exports.model_export import RefusedModelError

        refused = ModelResult(method="lasso", outcome="wellbeing",
                              refusal="34 case(s) are complete on every variable.")

        with pytest.raises(RefusedModelError, match="34 case"):
            build_model_docx(refused, request(), tmp_path / "m.docx")

        assert not (tmp_path / "m.docx").exists()

    def test_the_document_opens_without_a_matrix(self, fitted, tmp_path):
        """The matrix is optional; the caveats it carries are then simply
        absent rather than the export failing."""
        result, spec, _built = fitted

        assert build_model_docx(result, spec, tmp_path / "m.docx").exists()


class TestTheEffectLeadsWhenThereIsOne:

    def test_double_lasso_prints_its_interval(self, tmp_path):
        built = matrix(n=400)
        spec = request("double_lasso", treatment=("code::Cost",))
        result = fit(built, spec)

        text = docx_text(build_model_docx(result, spec, tmp_path / "m.docx",
                                          matrix=built))

        assert "Effect of code::Cost" in text
        assert "95% interval" in text

    def test_a_predictive_model_prints_no_interval(self, fitted, tmp_path):
        """Nothing but double lasso may show one, and an export inventing a
        column would be worse than omitting it."""
        result, spec, built = fitted

        text = docx_text(build_model_docx(result, spec, tmp_path / "m.docx",
                                          matrix=built))

        assert "95% interval" not in text
        assert "standard error" not in text.lower()

    def test_nothing_is_called_causal(self, tmp_path):
        built = matrix(n=400)
        spec = request("double_lasso", treatment=("code::Cost",))

        text = docx_text(build_model_docx(fit(built, spec), spec,
                                          tmp_path / "m.docx", matrix=built))

        assert "not a causal claim" in text


class TestTheWorkbook:

    def test_it_has_the_three_sheets(self, fitted, tmp_path):
        result, spec, built = fitted

        path = build_model_workbook(result, spec, tmp_path / "m.xlsx",
                                    matrix=built)
        sheets = pd.read_excel(path, sheet_name=None)

        assert set(sheets) == {SHEET_ESTIMATES, SHEET_SPEC, SHEET_CAVEATS}

    def test_the_caveats_are_a_sheet_of_their_own(self, fitted, tmp_path):
        """Merged into the estimates sheet they would be scrolled past."""
        result, spec, built = fitted

        path = build_model_workbook(result, spec, tmp_path / "m.xlsx",
                                    matrix=built)
        caveats = pd.read_excel(path, sheet_name=SHEET_CAVEATS)

        assert len(caveats) >= len(result.notes)

    def test_the_workbook_carries_no_response_text(self, fitted, tmp_path):
        """Every workbook in this program stays email-safe."""
        result, spec, built = fitted

        path = build_model_workbook(result, spec, tmp_path / "m.xlsx",
                                    matrix=built)
        blob = " ".join(
            str(value)
            for sheet in pd.read_excel(path, sheet_name=None).values()
            for value in sheet.to_numpy().ravel())

        assert "raw_text" not in blob

    def test_a_refused_model_is_not_written(self, tmp_path):
        """Found by a fixture that accidentally built an invalid spec: the
        export wrote a document with a title, a seed and nothing else."""
        from analysis.models import ModelResult

        from exports.model_export import RefusedModelError

        refused = ModelResult(method="lasso", outcome="wellbeing",
                              refusal="Not enough complete cases.")

        with pytest.raises(RefusedModelError, match="nothing to export"):
            build_model_workbook(refused, request(), tmp_path / "m.xlsx")

        assert not (tmp_path / "m.xlsx").exists()


class TestTheDesignMatrixLeavesWithItsSidecar:
    """Refusing to let an analyst take their data to R would be the mistake the
    REFI-QDA work exists to avoid. A matrix without the sidecar is a table of
    numbers nobody can reconstruct."""

    def test_the_csv_holds_every_case(self, tmp_path):
        built = matrix(n=120)

        path = export_design_matrix(built, tmp_path / "m.csv")

        assert len(pd.read_csv(path)) == 120

    def test_the_sidecar_is_written_beside_it(self, tmp_path):
        built = matrix()

        path = export_design_matrix(built, tmp_path / "m.csv")

        assert path.with_suffix(".json").exists()

    def test_the_sidecar_records_the_unit_and_the_coder(self, tmp_path):
        """Both change the answer, so a CSV without them is unreadable."""
        built = matrix()

        export_design_matrix(built, tmp_path / "m.csv")
        sidecar = json.loads((tmp_path / "m.json").read_text(encoding="utf-8"))

        assert sidecar["unit"] == "source_id"
        assert sidecar["rater"] == "human_primary"
        assert sidecar["code_encoding"] == "binary"

    def test_the_sidecar_records_reference_levels_and_drops(self, tmp_path):
        built = matrix()

        export_design_matrix(built, tmp_path / "m.csv")
        sidecar = json.loads((tmp_path / "m.json").read_text(encoding="utf-8"))

        assert sidecar["reference_levels"]["program"] == "Arts"
        assert "code::Rare" in sidecar["dropped_columns"]

    def test_the_sidecar_explains_zero_versus_missing(self, tmp_path):
        """The one thing an outside analyst will otherwise get wrong."""
        sidecar = design_matrix_sidecar(matrix())

        blob = " ".join(sidecar["notes"])
        assert "real zero" in blob
        assert "Do not fill it with zero" in blob

    def test_the_sidecar_is_plain_json(self, tmp_path):
        sidecar = design_matrix_sidecar(matrix())

        assert json.loads(json.dumps(sidecar)) == sidecar

    def test_the_sidecar_can_be_turned_off(self, tmp_path):
        built = matrix()

        export_design_matrix(built, tmp_path / "m.csv", include_sidecar=False)

        assert not (tmp_path / "m.json").exists()


class TestTheMethodsAppendixCarriesTheModel:
    """The document a reviewer asks for. A fitted model that is not in here is
    a model nobody can check."""

    @pytest.fixture()
    def project(self, tmp_path):
        from core.model import Code, CodeApplication, Excerpt, Source
        from core.project import Project

        path = tmp_path / "p.qnw"
        with Project.create(path) as project:
            code = project.add_code(Code(code_id="", name="Cost"))
            for index in range(6):
                text = f"response {index}"
                source_id = project.add_source(Source(
                    source_id="", type="survey_response", raw_text=text,
                    label="Q1", metadata={"response_id": f"r{index}"}))
                excerpt = project.add_excerpt(Excerpt(
                    excerpt_id="", source_id=source_id, text=text,
                    char_start=0, char_end=len(text)))
                project.apply_code(CodeApplication(
                    application_id="", excerpt_id=excerpt, code_id=code,
                    rater_id="human_primary"))
        return path

    def test_a_model_reaches_the_appendix(self, project, fitted, tmp_path):
        from exports.methods_appendix import build_methods_appendix

        result, spec, _built = fitted

        path = build_methods_appendix(project, tmp_path / "a.docx",
                                      models=[(result, spec)])
        text = docx_text(path)

        assert "Statistical models" in text
        assert "code::Cost" in text

    def test_its_caveats_reach_the_appendix_too(self, project, fitted, tmp_path):
        from exports.methods_appendix import build_methods_appendix

        result, spec, _built = fitted

        text = docx_text(build_methods_appendix(
            project, tmp_path / "a.docx", models=[(result, spec)]))

        for note in result.notes:
            assert note[:60] in text

    def test_the_appendix_says_codes_are_constructed_variables(
            self, project, fitted, tmp_path):
        from exports.methods_appendix import build_methods_appendix

        result, spec, _built = fitted

        text = docx_text(build_methods_appendix(
            project, tmp_path / "a.docx", models=[(result, spec)]))

        assert "constructed" in text and "not quantities" in text

    def test_the_seed_is_in_the_appendix(self, project, fitted, tmp_path):
        from exports.methods_appendix import build_methods_appendix

        result, spec, _built = fitted

        text = docx_text(build_methods_appendix(
            project, tmp_path / "a.docx", models=[(result, spec)]))

        assert f"seed {result.seed}" in text

    def test_an_appendix_without_models_is_unchanged(self, project, tmp_path):
        """Models are optional; every project that has none must still export."""
        from exports.methods_appendix import build_methods_appendix

        text = docx_text(build_methods_appendix(project, tmp_path / "a.docx"))

        assert "Statistical models" not in text
        assert "Inter-rater reliability" in text

    def test_models_are_placed_after_reliability(self, project, fitted, tmp_path):
        """A model shown before its coding's reliability invites reading the
        estimate as though the variables had been measured."""
        from exports.methods_appendix import build_methods_appendix

        result, spec, _built = fitted

        text = docx_text(build_methods_appendix(
            project, tmp_path / "a.docx", models=[(result, spec)]))

        assert text.index("Inter-rater reliability") < text.index("Statistical models")
        assert text.index("Statistical models") < text.index("Data handling")
