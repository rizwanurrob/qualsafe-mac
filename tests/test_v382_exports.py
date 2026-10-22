"""
Tests for v3.8.2 (P2) — Word report, Qualtrics ingestion, Excel workbook, .sav.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.ingest import strip_qualtrics_metadata
from core.model import Code, CodeApplication, Descriptor, Memo, Rater, Source
from core.project import Project, new_id


SECRET_QUOTE = "my advisor Dr Chen personally helped me every week"


def _fixture_store(path):
    """A store with codes, human+model codings, a memo quoting raw text."""
    with Project.create(path) as p:
        p.add_rater(Rater(rater_id="transformer", kind="transformer", name="model"))
        p.add_rater(Rater(rater_id="human_primary", kind="human", name="analyst"))
        belong = p.add_code(Code(code_id="", name="Belonging", definition="fitting in"))
        cost = p.add_code(Code(code_id="", name="Cost", definition="money worries"))
        rows = [("r1", SECRET_QUOTE, [belong]), ("r2", "tuition is too high", [cost]),
                ("r3", "I felt at home", [belong])]
        for rid, text, codes in rows:
            src = Source(source_id=new_id("src"), type="survey_response", raw_text=text,
                         label="Q1", metadata={"response_id": rid, "question_text": "Q1"})
            p.add_source(src)
            exc = p.add_whole_response_excerpt(src)
            for c in codes:
                p.apply_code(CodeApplication(application_id="", excerpt_id=exc, code_id=c,
                                             coder="transformer", rater_id="transformer"))
                p.apply_code(CodeApplication(application_id="", excerpt_id=exc, code_id=c,
                                             coder="human", rater_id="human_primary"))
        p.add_memo(Memo(memo_id="", target_type="project", target_id=None,
                        text=f"Striking quote: {SECRET_QUOTE}", author="human"))
    return path


# ── 8.8 Word report ──────────────────────────────────────────────────────────

def _docx_text(path):
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)


class TestDocxReport:
    def test_builds_and_has_core_sections(self, tmp_path):
        from exports.docx_report import build_docx_report

        store = _fixture_store(tmp_path / "s.qnw")
        out = build_docx_report(store, tmp_path / "report.docx")
        assert out.exists()
        text = _docx_text(out)
        assert "Codebook" in text and "Belonging" in text
        assert "frequency" in text.lower()

    def test_redacted_by_default_no_raw_text(self, tmp_path):
        from exports.docx_report import build_docx_report

        store = _fixture_store(tmp_path / "s2.qnw")
        out = build_docx_report(store, tmp_path / "r2.docx", include_raw_text=False)
        # Walk the actual docx XML, not just paragraph text, to be sure.
        with zipfile.ZipFile(out) as zf:
            blob = " ".join(zf.read(n).decode("utf-8", "replace")
                            for n in zf.namelist() if n.endswith(".xml"))
        assert SECRET_QUOTE not in blob
        assert "Dr Chen" not in blob

    def test_raw_text_included_when_asked(self, tmp_path):
        from exports.docx_report import build_docx_report

        store = _fixture_store(tmp_path / "s3.qnw")
        out = build_docx_report(store, tmp_path / "r3.docx", include_raw_text=True)
        assert SECRET_QUOTE in _docx_text(out)

    def test_agreement_section_when_both_raters(self, tmp_path):
        from exports.docx_report import build_docx_report

        store = _fixture_store(tmp_path / "s4.qnw")
        out = build_docx_report(store, tmp_path / "r4.docx")
        assert "agreement" in _docx_text(out).lower()


# ── 8.9 Qualtrics stripping ──────────────────────────────────────────────────

class TestQualtricsStrip:
    def test_removes_metadata_rows(self):
        df = pd.DataFrame({
            "response_id": ["Response ID", '{"ImportId":"_recordId"}', "1", "2"],
            "q1": ["What was hard?", '{"ImportId":"QID1"}', "cost", "belonging"],
        })
        out, n = strip_qualtrics_metadata(df)
        assert n == 2
        assert out["response_id"].tolist() == ["1", "2"]

    def test_no_op_on_normal_file(self):
        df = pd.DataFrame({"response_id": ["1", "2"], "text": ["a", "b"]})
        out, n = strip_qualtrics_metadata(df)
        assert n == 0
        assert len(out) == 2

    def test_importid_only_on_first_data_row(self):
        # A response that merely mentions ImportId shouldn't strip anything if it's
        # past the 3-row scan window.
        df = pd.DataFrame({"c": ["a", "b", "c", "row with ImportId text"]})
        out, n = strip_qualtrics_metadata(df)
        assert n == 0


# ── 8.10 Excel workbook ──────────────────────────────────────────────────────

class TestExcelWorkbook:
    def test_sheets_present_and_no_raw_text(self, tmp_path):
        from exports.excel_workbook import build_analysis_workbook

        store = _fixture_store(tmp_path / "x.qnw")
        out = build_analysis_workbook(store, tmp_path / "tables.xlsx")
        assert out.exists()
        xls = pd.ExcelFile(out)
        assert {"Codebook", "Frequency", "Co-occurrence"} <= set(xls.sheet_names)
        # No response text anywhere in the workbook.
        blob = ""
        for sheet in xls.sheet_names:
            blob += pd.read_excel(out, sheet_name=sheet).to_csv()
        assert SECRET_QUOTE not in blob

    def test_agreement_sheet_when_both_raters(self, tmp_path):
        from exports.excel_workbook import build_analysis_workbook

        store = _fixture_store(tmp_path / "x2.qnw")
        out = build_analysis_workbook(store, tmp_path / "t2.xlsx")
        assert "Agreement" in pd.ExcelFile(out).sheet_names


# ── 8.11 .sav graceful degradation ───────────────────────────────────────────

class TestSavSupport:
    def test_missing_pyreadstat_gives_actionable_message(self, tmp_path, monkeypatch):
        import streamlit_app as app

        # Simulate pyreadstat not installed.
        import builtins
        real_import = builtins.__import__

        def fake_import(name, *a, **k):
            if name == "pyreadstat":
                raise ImportError("no module")
            return real_import(name, *a, **k)

        monkeypatch.setattr(builtins, "__import__", fake_import)
        with pytest.raises(ValueError, match="pyreadstat"):
            app._read_sav(tmp_path / "x.sav")
