"""
Results & Reports: that each artefact is produced, and what is inside it.

The point of this screen is that `exports/` was already correct and simply
unreachable. So these tests spend most of their effort on the two things
wiring can get wrong: **the redaction contract** -- whether verbatim response
text ends up in a file -- and **the threading rule**, that a worker is handed a
path rather than the GUI thread's sqlite connection.

Every report is generated for real against a small project. Mocking the
builders would test only that buttons call functions, which is not the part
that breaks.
"""

from __future__ import annotations

import os
import zipfile

import pytest

pytest.importorskip("PySide6", reason="Qt front end not installed")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtGui import QUndoStack  # noqa: E402
from PySide6.QtWidgets import QApplication, QPushButton  # noqa: E402

from core.documents import add_document  # noqa: E402
from core.model import (  # noqa: E402
    Code, CodeApplication, Descriptor, Excerpt, Rater,
)
from core.project import Project  # noqa: E402
from qualsafe_qt.services.project_service import ProjectService  # noqa: E402
from qualsafe_qt.views.results_view import (  # noqa: E402
    REPORT_KINDS, ResultsView,
)

#: A distinctive phrase, so "did response text reach this file" is answerable
#: by searching for one string rather than by eyeballing.
SECRET = "the tuition was the hard part"
DOC_TEXT = (
    f"Respondent: {SECRET}, and the books on top of it. "
    "Making friends took most of the first term. "
    "My advisor was hard to reach in October."
)


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def coded_project(tmp_path):
    """Two human coders who agree on one code and differ on another."""
    path = tmp_path / "reports.qnw"
    with Project.create(path) as proj:
        source_id = add_document(proj, "Interview 01", DOC_TEXT)
        codes = {name: proj.add_code(Code(code_id="", name=name))
                 for name in ("Financial concerns", "Sense of belonging",
                              "Advising")}
        proj.add_rater(Rater(rater_id="human_primary", kind="human", name="A"))
        proj.add_rater(Rater(rater_id="human_secondary", kind="human", name="B"))

        spans = [(12, 60), (62, 105)]
        excerpts = [
            proj.add_excerpt(Excerpt(
                excerpt_id="", source_id=source_id, text=DOC_TEXT[a:b],
                char_start=a, char_end=b))
            for a, b in spans
        ]

        def apply(excerpt_id, code_name, rater):
            proj.apply_code(CodeApplication(
                application_id="", excerpt_id=excerpt_id,
                code_id=codes[code_name], coder="human", rater_id=rater))

        apply(excerpts[0], "Financial concerns", "human_primary")
        apply(excerpts[0], "Financial concerns", "human_secondary")
        apply(excerpts[1], "Sense of belonging", "human_primary")
        apply(excerpts[1], "Advising", "human_secondary")

        proj.set_descriptor(Descriptor(source_id=source_id, key="Program",
                                       value="Nursing"))
    return path


@pytest.fixture
def view(qapp, coded_project):
    service = ProjectService()
    service.open(coded_project)
    widget = ResultsView(service, QUndoStack())
    yield widget, service, coded_project
    service.close()


# -- every artefact is actually produced -------------------------------------


@pytest.mark.parametrize("kind", REPORT_KINDS, ids=lambda k: k.key)
def test_every_report_writes_a_non_empty_file(kind, coded_project, tmp_path):
    """Generated for real, not mocked. These builders already worked; the
    thing being tested is that this screen calls them correctly."""
    out = tmp_path / f"out{kind.suffix}"
    written = kind.builder(coded_project, out)

    assert written.exists(), f"{kind.key} reported success but wrote nothing"
    assert written.stat().st_size > 0
    assert written.suffix == kind.suffix


def test_the_methods_appendix_is_offered_first() -> None:
    """It is the one that makes a run reportable, so it leads the screen."""
    assert REPORT_KINDS[0].key == "methods"


# -- the redaction contract --------------------------------------------------


def _zip_text(path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", "replace")
            for name in archive.namelist()
            if not name.endswith((".png", ".xlsx"))
        )


def test_the_bundle_carries_no_response_text_by_default(coded_project, tmp_path):
    """The default must be safe. A researcher who clicks straight through
    should not discover later that quotes left the building."""
    from exports.bundle import build_bundle

    out = build_bundle(str(coded_project), tmp_path / "b.zip")
    assert SECRET not in _zip_text(out)


def test_the_bundle_can_carry_quotes_when_asked(coded_project, tmp_path):
    """The opt-in has to actually work, or the checkbox is a lie."""
    from exports.bundle import build_bundle

    out = build_bundle(str(coded_project), tmp_path / "b2.zip",
                       include_raw_text=True)
    assert SECRET in _zip_text(out)


def test_the_workbook_never_carries_response_text(coded_project, tmp_path):
    """Stated on the screen as "safe to email". That claim is load-bearing and
    the workbook has no option to change it."""
    from exports.excel_workbook import build_analysis_workbook

    out = build_analysis_workbook(str(coded_project), tmp_path / "w.xlsx")
    with zipfile.ZipFile(out) as archive:  # .xlsx is a zip of XML
        blob = "\n".join(
            archive.read(name).decode("utf-8", "replace")
            for name in archive.namelist() if name.endswith(".xml")
        )
    assert SECRET not in blob

    workbook_kind = next(k for k in REPORT_KINDS if k.key == "workbook")
    assert "never" in workbook_kind.contains.lower()


def test_every_report_states_what_it_contains() -> None:
    """Redaction differs per report, so each one says so beside its button
    rather than leaving the user to infer a single global policy."""
    for kind in REPORT_KINDS:
        assert kind.contains.strip(), f"{kind.key} does not say what it contains"


def test_options_that_add_quotes_are_off_by_default(view) -> None:
    widget, _service, _path = view
    assert not widget.report_raw_text.isChecked()
    assert not widget.bundle_raw_text.isChecked()
    assert not widget.codebook_examples.isChecked()


# -- the screen --------------------------------------------------------------


def test_nothing_can_be_written_without_a_project(qapp) -> None:
    service = ProjectService()
    widget = ResultsView(service, QUndoStack())
    assert not service.is_open
    assert widget._buttons, "expected a button per report"
    for key, button in widget._buttons.items():
        assert not button.isEnabled(), f"{key} is clickable with no project"


def test_opening_a_project_enables_the_buttons(view) -> None:
    widget, _service, _path = view
    for key, button in widget._buttons.items():
        assert button.isEnabled(), f"{key} stayed disabled with a project open"


def test_there_is_a_button_for_every_report_kind(view) -> None:
    widget, _service, _path = view
    assert set(widget._buttons) == {k.key for k in REPORT_KINDS}
    assert all(isinstance(b, QPushButton) for b in widget._buttons.values())


def test_the_worker_is_given_a_path_not_the_open_project(view, tmp_path,
                                                         monkeypatch) -> None:
    """A sqlite connection belongs to the thread that opened it. Handing the
    GUI thread's `Project` to a worker is the failure this rule exists to
    prevent, and it would not show up as an exception every time -- it can
    corrupt instead. So assert the argument is a path.
    """
    widget, _service, store = view
    captured = {}

    def fake_run_task(fn, *args, **kwargs):
        captured["fn"] = fn
        captured["args"] = args
        captured["reads_store"] = kwargs.get("reads_store")
        return None

    monkeypatch.setattr("qualsafe_qt.views.results_view.run_task", fake_run_task)
    monkeypatch.setattr(
        "qualsafe_qt.views.results_view.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: (str(tmp_path / "x.docx"), "")))

    kind = next(k for k in REPORT_KINDS if k.key == "methods")
    widget._save(kind)

    assert captured, "the save never reached run_task"
    first = captured["args"][0]
    assert not hasattr(first, "coding_table"), (
        "an open Project was handed to a worker thread; pass the .qnw path")
    assert str(first).endswith(".qnw")
    assert captured["reads_store"] is True, (
        "reads_store=True is what tells the caller to reopen the GUI "
        "connection after the worker touched the store")


def test_a_cancelled_save_dialog_writes_nothing(view, monkeypatch) -> None:
    widget, _service, _path = view
    called = {"ran": False}

    def fake_run_task(*_a, **_k):
        called["ran"] = True

    monkeypatch.setattr("qualsafe_qt.views.results_view.run_task", fake_run_task)
    monkeypatch.setattr(
        "qualsafe_qt.views.results_view.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: ("", "")))

    widget._save(REPORT_KINDS[0])
    assert not called["ran"], "cancelling the dialog still started a job"


def test_including_quotes_requires_confirmation(view, tmp_path,
                                                monkeypatch) -> None:
    """Verbatim text leaving the project is the one irreversible thing this
    screen does, so it is confirmed once, plainly, before the file is written."""
    widget, _service, _path = view
    widget.report_raw_text.setChecked(True)
    started = {"ran": False}

    monkeypatch.setattr("qualsafe_qt.views.results_view.run_task",
                        lambda *a, **k: started.update(ran=True))
    monkeypatch.setattr(
        "qualsafe_qt.views.results_view.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: (str(tmp_path / "r.docx"), "")))
    monkeypatch.setattr(widget, "_confirm_raw_text", lambda _kind: False)

    widget._save(next(k for k in REPORT_KINDS if k.key == "report"))
    assert not started["ran"], "declining the warning still wrote the file"

    monkeypatch.setattr(widget, "_confirm_raw_text", lambda _kind: True)
    widget._save(next(k for k in REPORT_KINDS if k.key == "report"))
    assert started["ran"], "accepting the warning did not write the file"


def test_group_titles_escape_the_ampersand() -> None:
    """QGroupBox reads a single `&` as a keyboard mnemonic, so "Methods &
    reliability" would render as "Methods _reliability"."""
    for kind in REPORT_KINDS:
        if "&" in kind.title:
            assert "&&" in kind.title, (
                f"{kind.key} title needs && to render a literal ampersand")


# -- the joint display -------------------------------------------------------
#
# The mixed-methods deliverable. Word is the table a findings chapter prints;
# Excel is the working file, and like every other workbook here it must never
# carry response text.


def _joint_kind():
    return next(k for k in REPORT_KINDS if k.key == "joint")


def test_the_joint_display_offers_the_projects_descriptors(view) -> None:
    widget, _service, _path = view
    options = [widget.joint_descriptor.itemText(i)
               for i in range(widget.joint_descriptor.count())]
    assert options == ["Program"]


def test_the_joint_display_is_refused_without_a_descriptor(qapp, tmp_path) -> None:
    """There is no subgroup axis to put across the top, so the button says so
    rather than writing an empty table."""
    path = tmp_path / "bare.qnw"
    Project.create(path).close()
    service = ProjectService()
    service.open(path)
    widget = ResultsView(service, QUndoStack())

    assert not widget._buttons["joint"].isEnabled()
    assert "no respondent details" in widget._buttons["joint"].toolTip()
    service.close()


def test_the_joint_display_workbook_carries_no_response_text(
        coded_project, tmp_path) -> None:
    import pandas as pd

    out = _joint_kind().builder(coded_project, tmp_path / "joint.xlsx",
                                descriptor_key="Program")

    sheets = pd.read_excel(out, sheet_name=None)
    blob = " ".join(frame.to_string() for frame in sheets.values())
    assert SECRET not in blob
    assert {"Joint display", "Subgroup sizes", "Caveats"} == set(sheets)


def test_quotes_in_the_word_table_are_opt_in(coded_project, tmp_path) -> None:
    from docx import Document

    def cells(path):
        return " ".join(c.text for t in Document(path).tables
                        for r in t.rows for c in r.cells)

    with_quotes = _joint_kind().builder(
        coded_project, tmp_path / "with.docx", descriptor_key="Program",
        include_quotes=True)
    without = _joint_kind().builder(
        coded_project, tmp_path / "without.docx", descriptor_key="Program",
        include_quotes=False)

    assert SECRET in cells(with_quotes)
    assert SECRET not in cells(without)


def test_the_integration_statements_written_on_the_analysis_screen_travel(
        coded_project, tmp_path) -> None:
    """They are memos in the project, so both front ends read the same ones."""
    from docx import Document

    from core.model import Memo

    with Project.open(coded_project) as proj:
        proj.add_memo(Memo(memo_id="", target_type="joint_display",
                           target_id="Program::Financial concerns",
                           text="Nursing students carry this alone."))

    out = _joint_kind().builder(coded_project, tmp_path / "joint.docx",
                                descriptor_key="Program")

    cells = " ".join(c.text for t in Document(out).tables
                     for r in t.rows for c in r.cells)
    assert "Nursing students carry this alone." in cells
