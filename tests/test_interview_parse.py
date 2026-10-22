"""
Tests for the v3.1 interview parsing core (I1): template schema, document
loaders, turn segmentation (all fixture formats + fallback), speaker-role
inference, question detection, and Q/A pairing.

Offline: fixtures are synthetic .txt transcripts; the DOCX loader is exercised
against a file generated at test time with python-docx; the PDF loader against a
mocked pdfplumber (no binary fixtures in the repo).
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from interview.parse import Turn, load_document, parse_interview, segment_turns
from interview.questions import detect_questions, pair_questions_with_responses
from interview.speakers import apply_roles, infer_speaker_roles
from interview.template import InterviewTemplate, TemplateError

FIXTURES = Path(__file__).resolve().parent / "fixtures"
EXAMPLE_TEMPLATE = Path(__file__).resolve().parent.parent / "config" / "interview_template.example.yaml"


def _fixture_text(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


# ─────────────────────────────────────────────────────────────────────────────
# Template schema
# ─────────────────────────────────────────────────────────────────────────────

class TestTemplate:
    def test_loads_example_template(self):
        t = InterviewTemplate.load(EXAMPLE_TEMPLATE)
        assert t.title == "Student Belonging Interview Guide"
        # v3.7.1 extended the example with finances + improvement so it pairs
        # with the built-in sample transcript.
        assert [th.id for th in t.themes] == [
            "transition", "belonging", "support", "finances", "improvement"]
        assert "belong" in t.themes[1].keywords

    def test_missing_themes_raises(self):
        with pytest.raises(TemplateError, match="themes"):
            InterviewTemplate.from_dict({"interview_template": {"title": "x"}})

    def test_theme_missing_label_raises(self):
        data = {"interview_template": {"themes": [{"id": "a"}]}}
        with pytest.raises(TemplateError, match="label"):
            InterviewTemplate.from_dict(data)

    def test_duplicate_ids_raise(self):
        data = {"interview_template": {"themes": [
            {"id": "a", "label": "A"}, {"id": "a", "label": "A2"}]}}
        with pytest.raises(TemplateError, match="Duplicate"):
            InterviewTemplate.from_dict(data)

    def test_to_codes_maps_apriori_codebook(self):
        t = InterviewTemplate.load(EXAMPLE_TEMPLATE)
        codes = t.to_codes()
        assert [c.code_id for c in codes] == [
            "transition", "belonging", "support", "finances", "improvement"]
        assert all(c.origin == "a_priori" and c.created_by == "human" for c in codes)
        assert codes[0].definition.startswith("Adjustment")


# ─────────────────────────────────────────────────────────────────────────────
# Turn segmentation — one class per fixture format + shared guarantees
# ─────────────────────────────────────────────────────────────────────────────

class TestSegmentation:
    @pytest.mark.parametrize("fixture,expected_pattern,expected_turns,expected_labels", [
        ("interviewer_respondent.txt", "interviewer_respondent", 6, {"Interviewer", "Respondent"}),
        ("q_a.txt", "q_a", 6, {"Q", "A"}),
        ("initials.txt", "name_colon", 6, {"MK", "JT"}),
        ("timestamped.txt", "timestamped", 6, {"Speaker 1", "Speaker 2"}),
    ])
    def test_formats_detected(self, fixture, expected_pattern, expected_turns, expected_labels):
        raw = _fixture_text(fixture)
        turns, pattern, confidence = segment_turns(raw)
        assert pattern == expected_pattern
        assert len(turns) == expected_turns
        assert {t.speaker_raw for t in turns} == expected_labels
        assert confidence > 0.5

    def test_offsets_verbatim_for_all_fixtures(self):
        """Every turn's text must equal raw[start:end] — the traceability guarantee."""
        for name in ["interviewer_respondent.txt", "q_a.txt", "initials.txt", "timestamped.txt"]:
            raw = _fixture_text(name)
            turns, _, _ = segment_turns(raw)
            for t in turns:
                assert raw[t.char_start:t.char_end] == t.text, f"offset drift in {name} turn {t.index}"

    def test_multiline_turn_includes_continuation(self):
        raw = _fixture_text("interviewer_respondent.txt")
        turns, _, _ = segment_turns(raw)
        # Respondent's second turn continues onto an unlabeled line.
        second_resp = [t for t in turns if t.speaker_raw == "Respondent"][1]
        assert "That made a big difference" in second_resp.text

    def test_messy_document_falls_back_unguessed(self):
        raw = _fixture_text("messy.txt")
        turns, pattern, confidence = segment_turns(raw)
        assert pattern == "none"
        assert confidence == 0.0
        assert len(turns) == 1 and turns[0].speaker_raw == "UNKNOWN"

    def test_forced_pattern(self):
        raw = _fixture_text("q_a.txt")
        turns, pattern, _ = segment_turns(raw, pattern_name="q_a")
        assert pattern == "q_a" and len(turns) == 6

    def test_unknown_pattern_name_raises(self):
        with pytest.raises(ValueError, match="Unknown speaker pattern"):
            segment_turns("x", pattern_name="nope")


# ─────────────────────────────────────────────────────────────────────────────
# Document loaders
# ─────────────────────────────────────────────────────────────────────────────

class TestLoaders:
    def test_txt(self, tmp_path):
        p = tmp_path / "t.txt"
        p.write_text("Interviewer: hi\nRespondent: hello", encoding="utf-8")
        assert "Respondent: hello" in load_document(p)

    def test_docx_generated_at_test_time(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Interviewer: How was your week?")
        doc.add_paragraph("Respondent: Busy, but good.")
        p = tmp_path / "t.docx"
        doc.save(str(p))
        raw = load_document(p)
        turns, pattern, _ = segment_turns(raw)
        assert pattern == "interviewer_respondent" and len(turns) == 2

    def test_pdf_via_mocked_pdfplumber(self, tmp_path):
        p = tmp_path / "t.pdf"
        p.write_bytes(b"%PDF-fake")  # existence check only; pdfplumber is mocked
        page1, page2 = MagicMock(), MagicMock()
        page1.extract_text.return_value = "Q: How are classes?\nA: Fine so far."
        page2.extract_text.return_value = "Q: What is hardest?\nA: The commute."
        pdf = MagicMock()
        pdf.pages = [page1, page2]
        cm = MagicMock()
        cm.__enter__ = MagicMock(return_value=pdf)
        cm.__exit__ = MagicMock(return_value=False)
        with patch("pdfplumber.open", return_value=cm):
            raw = load_document(p)
        turns, pattern, _ = segment_turns(raw)
        assert pattern == "q_a" and len(turns) == 4

    def test_unsupported_extension_raises(self, tmp_path):
        p = tmp_path / "t.rtf"
        p.write_text("x", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            load_document(p)

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            load_document(tmp_path / "nope.txt")


# ─────────────────────────────────────────────────────────────────────────────
# Speaker roles
# ─────────────────────────────────────────────────────────────────────────────

class TestSpeakerRoles:
    @pytest.mark.parametrize("fixture,interviewer_label", [
        ("interviewer_respondent.txt", "Interviewer"),
        ("q_a.txt", "Q"),
        ("initials.txt", "MK"),
        ("timestamped.txt", "Speaker 1"),
    ])
    def test_interviewer_inferred(self, fixture, interviewer_label):
        turns, _, _ = segment_turns(_fixture_text(fixture))
        inference = infer_speaker_roles(turns)
        assert inference.role_by_label[interviewer_label] == "interviewer"
        assert inference.confidence > 0.0
        assert interviewer_label in inference.rationale

    def test_single_speaker_abstains(self):
        turns = [Turn(0, "A", "hello there", 0, 11)]
        inference = infer_speaker_roles(turns)
        assert inference.confidence == 0.0
        assert "review" in inference.rationale.lower()

    def test_apply_roles_and_third_speaker_is_other(self):
        turns, _, _ = segment_turns(_fixture_text("q_a.txt"))
        roled = apply_roles(turns, {"Q": "interviewer", "A": "interviewee"})
        assert all(t.speaker_role in {"interviewer", "interviewee"} for t in roled)
        extra = apply_roles([Turn(0, "X", "hm", 0, 2)], {"Q": "interviewer"})
        assert extra[0].speaker_role == "other"


# ─────────────────────────────────────────────────────────────────────────────
# Question detection + Q/A pairing
# ─────────────────────────────────────────────────────────────────────────────

def _roled(fixture: str):
    turns, _, _ = segment_turns(_fixture_text(fixture))
    inference = infer_speaker_roles(turns)
    return apply_roles(turns, inference.role_by_label)


class TestQuestions:
    @pytest.mark.parametrize("fixture,expected_n", [
        ("interviewer_respondent.txt", 3),
        ("q_a.txt", 3),
        ("initials.txt", 3),   # includes implicit "Tell me about..." prompt
        ("timestamped.txt", 3),
    ])
    def test_detection_counts(self, fixture, expected_n):
        qs = detect_questions(_roled(fixture))
        assert len(qs) == expected_n

    def test_implicit_prompt_detected(self):
        qs = detect_questions(_roled("initials.txt"))
        assert any(q.text.startswith("Tell me about") for q in qs)

    def test_backchannel_not_a_question(self):
        turns = [Turn(0, "I", "Right?", 0, 6, speaker_role="interviewer")]
        assert detect_questions(turns) == []

    def test_interviewee_questions_ignored(self):
        turns = [Turn(0, "R", "What do you mean by that?", 0, 25, speaker_role="interviewee")]
        assert detect_questions(turns) == []

    def test_pairing_attaches_following_response(self):
        turns = _roled("q_a.txt")
        qs = detect_questions(turns)
        pairs = pair_questions_with_responses(turns, qs)
        assert len(pairs) == 3
        assert all(len(p.response_turns) == 1 for p in pairs)
        assert "management felt natural" in pairs[0].response_turns[0].text

    def test_same_turn_followup_attaches_to_last_question(self):
        turns = [
            Turn(0, "I", "How was the fall? What surprised you most?", 0, 43, speaker_role="interviewer"),
            Turn(1, "R", "The workload surprised me the most.", 44, 80, speaker_role="interviewee"),
        ]
        qs = detect_questions(turns)
        assert len(qs) == 2
        pairs = pair_questions_with_responses(turns, qs)
        assert pairs[0].response_turns == []            # earlier question: no direct block
        assert len(pairs[1].response_turns) == 1        # last question takes the answer

    def test_end_to_end_parse_interview(self):
        raw, turns, pattern, confidence = parse_interview(FIXTURES / "interviewer_respondent.txt")
        assert pattern == "interviewer_respondent"
        assert confidence > 0.5
        assert raw[turns[0].char_start:turns[0].char_end] == turns[0].text
