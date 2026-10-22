#!/usr/bin/env python
"""
Codebook report — the appendix a dissertation or an IR deliverable needs
(v4, Phase 5.1).

A formatted Word document of the codebook as it actually stands: the hierarchy,
each code's definition and origin, how often it was applied, and a few example
passages. The codebook has lived only inside the app until now; this is the
artifact you attach to a report or hand to a reviewer.

Follows the same redaction rule as every other export: with
``include_examples=False`` (the default) no verbatim response text appears, so
the document is safe to circulate.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from core.manifest import PACKAGE_VERSION


def _depth(code, by_id: Dict[str, Any]) -> int:
    """How deep a code sits in the hierarchy (root = 0), cycle-safe."""
    depth, seen, cur = 0, {code.code_id}, code.parent_id
    while cur and cur in by_id and cur not in seen:
        depth += 1
        seen.add(cur)
        cur = by_id[cur].parent_id
    return depth


def codebook_rows(project) -> List[Dict[str, Any]]:
    """The codebook flattened for reporting: hierarchy order, depth, counts.

    Pure over the store — the same rows back the Word report and any other
    rendering. Children follow their parent, siblings alphabetically, so the
    order is stable and readable.
    """
    codes = project.get_codebook()
    by_id = {c.code_id: c for c in codes}
    freq = project.code_frequency()
    children: Dict[Optional[str], List[Any]] = {}
    for code in codes:
        parent = code.parent_id if code.parent_id in by_id else None
        children.setdefault(parent, []).append(code)
    for group in children.values():
        group.sort(key=lambda c: c.name.lower())

    rows: List[Dict[str, Any]] = []

    def walk(parent: Optional[str]) -> None:
        for code in children.get(parent, []):
            rows.append({
                "code_id": code.code_id,
                "name": code.name,
                "definition": code.definition or "",
                "origin": code.origin,
                "created_by": code.created_by,
                "depth": _depth(code, by_id),
                "parent": by_id[code.parent_id].name if code.parent_id in by_id else "",
                "applications": int(freq.get(code.code_id, 0)),
            })
            walk(code.code_id)

    walk(None)
    return rows


def _examples(project, code_id: str, limit: int) -> List[str]:
    """A few passages this code was applied to, shortest first (readable quotes)."""
    coding = project.coding_table()
    if coding.empty:
        return []
    excerpt_ids = coding[coding["code_id"] == code_id]["excerpt_id"].unique().tolist()
    texts = []
    for excerpt_id in excerpt_ids:
        excerpt = project.get_excerpt(excerpt_id)
        if excerpt is not None and excerpt.text.strip():
            texts.append(excerpt.text.strip())
    texts.sort(key=len)
    return texts[:limit]


def build_codebook_report(
    project,
    out_path: Union[str, Path],
    *,
    title: str = "Codebook",
    include_examples: bool = False,
    examples_per_code: int = 2,
    example_chars: int = 220,
) -> Path:
    """Write the codebook as a Word document. Returns the path.

    `project` may be an open Project or a path to a `.qnw`.
    """
    from docx import Document
    from docx.shared import Pt

    from core.project import Project

    close_after = False
    if not hasattr(project, "get_codebook"):
        project = Project.open(project)
        close_after = True

    out_path = Path(out_path)
    try:
        rows = codebook_rows(project)
        doc = Document()
        doc.add_heading(title, level=0)
        meta = doc.add_paragraph()
        run = meta.add_run(
            f"Generated {datetime.now().strftime('%B %d, %Y')} · QualSafe "
            f"{PACKAGE_VERSION} · {len(rows)} code(s) · "
            f"{'with' if include_examples else 'without'} example quotes")
        run.font.size = Pt(9)

        versions = project.list_codebook_versions()
        if versions:
            latest = versions[0]
            note = doc.add_paragraph()
            vrun = note.add_run(
                f"Codebook version {latest.get('codebook_hash', '')} "
                f"(last change: {latest.get('trigger') or 'n/a'} "
                f"on {latest.get('created_at', '')})")
            vrun.font.size = Pt(9)

        if not rows:
            doc.add_paragraph("This project has no codes.")
            doc.save(str(out_path))
            return out_path

        doc.add_heading("Codes", level=1)
        for row in rows:
            heading = doc.add_paragraph()
            heading.paragraph_format.left_indent = Pt(14 * row["depth"])
            name_run = heading.add_run(row["name"])
            name_run.bold = True
            tail = heading.add_run(
                f"   ({row['applications']} application"
                f"{'' if row['applications'] == 1 else 's'}, {row['origin']})")
            tail.font.size = Pt(9)

            if row["definition"]:
                body = doc.add_paragraph(row["definition"])
                body.paragraph_format.left_indent = Pt(14 * row["depth"] + 14)
            else:
                body = doc.add_paragraph("(no definition)")
                body.paragraph_format.left_indent = Pt(14 * row["depth"] + 14)
                body.runs[0].italic = True

            if include_examples and row["applications"]:
                for text in _examples(project, row["code_id"], examples_per_code):
                    quote = doc.add_paragraph(f"“{text[:example_chars]}”",
                                              style="Intense Quote")
                    quote.paragraph_format.left_indent = Pt(14 * row["depth"] + 20)

        doc.save(str(out_path))
        return out_path
    finally:
        if close_after:
            project.close()
