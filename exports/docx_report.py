#!/usr/bin/env python
"""
Word report export (v3.8.2, item 8.8).

Institutional research deliverables are Word documents; a markdown file asks the
analyst to do the last mile by hand. This builds a polished .docx straight from
the project store — the same source of truth as the shareable bundle, with the
same redaction rule: with ``include_raw_text=False`` (default) no verbatim
response text appears anywhere, and memos are scrubbed of any quoted source text.

Covers both modes with one report shape: codebook (with definitions), code
frequency, top co-occurring pairs, human-vs-model agreement when present, memos
(which is where interview per-theme LLM summaries live), and — only when raw text
is explicitly allowed — up to three example excerpts per code.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional, Union

from analysis.frequency import code_frequency
from exports.bundle import REDACTED, _redact_memo


def build_docx_report(
    project,
    out_path: Union[str, Path],
    *,
    title: str = "Qualitative Analysis Report",
    include_raw_text: bool = False,
    include_memos: bool = True,
    examples_per_code: int = 3,
) -> Path:
    """Build a Word report from a Project (or a path to a .qnw). Returns the path."""
    from docx import Document
    from docx.shared import Pt

    from core.project import Project

    close_after = False
    if not hasattr(project, "coding_table"):
        project = Project.open(project)
        close_after = True

    out_path = Path(out_path)
    try:
        doc = Document()
        doc.add_heading(title, level=0)
        meta = doc.add_paragraph()
        run = meta.add_run(
            f"Generated {datetime.now().strftime('%B %d, %Y')} · "
            f"QualSafe · all processing local · "
            f"verbatim quotes {'included' if include_raw_text else 'excluded (redacted)'}"
        )
        run.font.size = Pt(9)

        codes = project.get_codebook()
        coding = project.coding_table()
        name_by_id = {c.code_id: c.name for c in codes}

        # ── Codebook ────────────────────────────────────────────────────────
        doc.add_heading("Codebook", level=1)
        if codes:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "Code", "Definition", "Parent"
            for code in codes:
                row = table.add_row().cells
                row[0].text = code.name
                row[1].text = code.definition or ""
                row[2].text = name_by_id.get(code.parent_id, "")
        else:
            doc.add_paragraph("No codes in this project.")

        # ── Frequency ───────────────────────────────────────────────────────
        doc.add_heading("Code frequency", level=1)
        if not coding.empty:
            freq = code_frequency(coding)
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Code", "Applications"
            for _, r in freq.iterrows():
                row = table.add_row().cells
                row[0].text = str(r["code_name"])
                row[1].text = str(int(r["n"]))
        else:
            doc.add_paragraph("No code applications yet.")

        # ── Co-occurrence (top pairs) ───────────────────────────────────────
        if not coding.empty:
            from analysis.cooccurrence import cooccurrence_matrix

            matrix, _ = cooccurrence_matrix(coding)
            pairs = []
            cols = list(matrix.columns)
            for i, a in enumerate(cols):
                for b in cols[i + 1:]:
                    n = int(matrix.at[a, b])
                    if n > 0:
                        pairs.append((n, a, b))
            if pairs:
                doc.add_heading("Most frequent code co-occurrences", level=1)
                for n, a, b in sorted(pairs, reverse=True)[:10]:
                    doc.add_paragraph(f"{a}  +  {b}: {n}", style="List Bullet")

        # ── Agreement ───────────────────────────────────────────────────────
        agreement = _agreement_table(project)
        if agreement is not None:
            doc.add_heading("Human–model agreement (per-code Cohen's κ)", level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
                "Code", "κ", "% agreement", "n")
            for _, r in agreement.iterrows():
                row = table.add_row().cells
                row[0].text = str(r["code"])
                row[1].text = "insufficient N" if r["kappa"] != r["kappa"] else f"{r['kappa']:.2f}"
                row[2].text = f"{r['percent_agreement']:.0%}"
                row[3].text = str(int(r["n_items"]))
            doc.add_paragraph(f"Compared: {agreement.attrs.get('comparison', '')}.")
            for caveat in agreement.attrs.get("caveats", []):
                doc.add_paragraph(caveat, style="List Bullet")

        # ── Memos (incl. interview per-theme summaries) ─────────────────────
        if include_memos:
            memos = project.list_memos()
            if memos:
                doc.add_heading("Analytic memos", level=1)
                for memo in memos:
                    target = name_by_id.get(memo.target_id, memo.target_type)
                    body = memo.text if include_raw_text else _redact_memo(memo.text, project)
                    para = doc.add_paragraph()
                    para.add_run(f"[{target}] ").bold = True
                    para.add_run(body)

        # ── Example excerpts (raw text ONLY when explicitly allowed) ────────
        if include_raw_text and not coding.empty:
            doc.add_heading("Example excerpts by code", level=1)
            excerpt_by_id = {e.excerpt_id: e for e in project.list_excerpts()}
            freq = code_frequency(coding)
            for _, r in freq.iterrows():
                doc.add_heading(str(r["code_name"]), level=2)
                for excerpt_id in list(r["excerpt_ids"])[:examples_per_code]:
                    exc = excerpt_by_id.get(excerpt_id)
                    if exc is not None:
                        doc.add_paragraph(f"“{exc.text}”", style="Intense Quote")

        doc.save(str(out_path))
    finally:
        if close_after:
            project.close()
    return out_path


def _agreement_table(project):
    """Human-vs-engine kappa table (caveats in ``.attrs``), or None.

    The two raters are chosen by role in `core.agreement_pairs`, the same way
    for every export -- see that module for why alphabetical choice was wrong.
    """
    from core.agreement_pairs import human_vs_engine, kappa_table

    return kappa_table(human_vs_engine(project))
