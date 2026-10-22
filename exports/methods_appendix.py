#!/usr/bin/env python
"""
Methods & reliability appendix (v4, Phase 5.2).

The document a reviewer, an IRB, or a dissertation committee asks for: how the
coding was actually produced, by whom, how much the coders agreed, and what the
numbers do and do not support.

QualSafe already records everything needed — run manifests, rater identities,
per-code kappa, adjudication decisions, codebook versions, offline enforcement —
but it lived scattered across pages. This assembles it into one honest statement,
**including the caveats**. Reliability figures reported without their caveats are
how a kappa on six items ends up in a paper as though it meant something, so the
insufficient-N flags, the kappa-paradox note, and the unit of analysis travel
with the table rather than being left for the reader to remember.

Contains no verbatim response text, so it is safe to circulate.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd

from core.manifest import PACKAGE_VERSION, PRIVACY_NOTE

ADJUDICATED_RATER = "human_adjudicated"


def methods_facts(project) -> Dict[str, Any]:
    """Everything the appendix states, gathered from the store.

    Pure-ish (reads only), and separated from rendering so the facts can be
    tested without opening a Word document.
    """
    runs = project.list_runs()
    latest = runs[-1] if runs else None
    raters = project.list_raters()
    coding = project.coding_table()

    human_raters = [r for r in raters if r.kind == "human"]
    machine_raters = [r for r in raters if r.kind != "human"]
    adjudicated = 0
    if not coding.empty:
        adjudicated = int((coding["rater_id"] == ADJUDICATED_RATER).sum())

    versions = project.list_codebook_versions()
    from core.project import SCHEMA_VERSION

    return {
        "package_version": PACKAGE_VERSION,
        # The appendix is what a reviewer reads to reproduce a run, and the
        # package version alone does not tell them whether their install can
        # open the project file: a store from a newer schema raises rather than
        # degrading. Both numbers, or neither is actionable.
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "n_sources": project.count_sources(),
        "n_codes": len(project.get_codebook()),
        "n_codings": 0 if coding.empty else len(coding),
        "engine": getattr(latest, "engine", "") if latest else "",
        "model": getattr(latest, "model", "") if latest else "",
        "mode": getattr(latest, "mode", "") if latest else "",
        "thresholds": dict(getattr(latest, "thresholds", {}) or {}) if latest else {},
        "run_timestamp": getattr(latest, "timestamp", "") if latest else "",
        "n_runs": len(runs),
        "human_raters": [r.name or r.rater_id for r in human_raters],
        "machine_raters": [f"{r.rater_id} ({r.model_id})" if r.model_id else r.rater_id
                           for r in machine_raters],
        "adjudicated_codings": adjudicated,
        # Codings from the Alt Survey Flow demo engine (hash scores). They are
        # disclosed, never described as coding (review finding F2).
        "demo_codings": _demo_codings(project, coding),
        "codebook_version": versions[0].get("codebook_hash", "") if versions else "",
        "codebook_versions": len(versions),
    }


def _demo_codings(project, coding) -> int:
    from engines.laya_engine import demo_rater_ids

    demo = set(demo_rater_ids(project))
    if coding.empty or not demo:
        return 0
    return int(coding["rater_id"].isin(demo).sum())


def reliability_tables(project, *, min_support: int = 5) -> Dict[str, pd.DataFrame]:
    """Per-code kappa for each comparable rater pair.

    Returns ``{"human_primary vs rater_…": df, "human_A vs human_B": df, ...}``,
    each table carrying its own caveats in ``df.attrs["caveats"]``; empty when
    there is nothing comparable, which the appendix then says plainly instead of
    printing a table of NaNs. Pairs are chosen by `core.agreement_pairs`, and
    span-coded documents both coders finished are compared on sentences.
    """
    from core.agreement_pairs import comparisons, kappa_table, span_comparison, span_raters

    out: Dict[str, pd.DataFrame] = {}
    for comparison in comparisons(project):
        table = kappa_table(comparison, min_support=min_support)
        if table is not None:
            out[comparison.label] = table
    coders = span_raters(project)
    for i, a in enumerate(coders):
        for b in coders[i + 1:]:
            table = kappa_table(span_comparison(project, a, b, unit="sentence"),
                                min_support=min_support)
            if table is not None:
                out[table.attrs["comparison"]] = table
    return out


CAVEATS = [
    "Kappa is reported per code as binary Cohen's kappa over the items both "
    "raters reviewed, including items a rater reviewed and applied no code to. "
    "Codes below the minimum support are flagged insufficient-N "
    "and given no kappa rather than an unstable one.",
    "Percent agreement and prevalence are shown alongside kappa: for a rare code, "
    "high agreement with low kappa is the kappa paradox, not a coding failure.",
    "Agreement is computed on the unit recorded for the run. Figures computed at "
    "different units (whole response vs sentence vs span) are not comparable.",
    "Where a coding sample was selected by model uncertainty, it over-represents "
    "hard cases and therefore understates agreement on the corpus as a whole.",
]


def build_methods_appendix(
    project,
    out_path: Union[str, Path],
    *,
    title: str = "Methods and Reliability Appendix",
    min_support: int = 5,
    sampling_note: str = "",
    models: Optional[List[Tuple[Any, Any]]] = None,
) -> Path:
    """Write the appendix as a Word document. Returns the path.

    `models` is a list of `(result, request)` pairs from `analysis/models.fit`.
    A model reported in a paper needs its specification, its seed, its effective
    n and the provenance of every coded variable, and this is the document a
    reviewer asks for — so a fitted model that is not in here is a model nobody
    can check. `V6_PLAN.md` Phase D item 11.
    """
    from docx import Document
    from docx.shared import Pt

    from core.project import Project

    close_after = False
    if not hasattr(project, "coding_table"):
        project = Project.open(project)
        close_after = True

    out_path = Path(out_path)
    try:
        facts = methods_facts(project)
        tables = reliability_tables(project, min_support=min_support)

        doc = Document()
        doc.add_heading(title, level=0)
        stamp = doc.add_paragraph()
        run = stamp.add_run(
            f"Generated {datetime.now().strftime('%B %d, %Y')} · "
            f"QualSafe {facts['package_version']} · "
            f"project format v{facts['schema_version']}")
        run.font.size = Pt(9)

        # ── how the coding was produced ──
        doc.add_heading("Analysis procedure", level=1)
        if facts["engine"]:
            doc.add_paragraph(
                f"Coding was produced in {facts['mode'] or 'survey'} mode using the "
                f"{facts['engine']} engine with model {facts['model'] or 'n/a'} "
                f"({facts['n_runs']} run(s); most recent {facts['run_timestamp']}).")
            if facts["thresholds"]:
                pairs = ", ".join(f"{k} = {v}" for k, v in facts["thresholds"].items()
                                  if v is not None)
                if pairs:
                    doc.add_paragraph(
                        f"Automated codings were gated by confidence thresholds "
                        f"({pairs}); responses below the gate were routed to human "
                        f"review rather than coded.")
        else:
            doc.add_paragraph("No automated run is recorded in this project; the "
                              "coding below is human-applied.")
        doc.add_paragraph(
            f"The project contains {facts['n_sources']} source(s), "
            f"{facts['n_codes']} code(s), and {facts['n_codings']} coding(s).")
        if sampling_note.strip():
            doc.add_heading("Sampling", level=2)
            doc.add_paragraph(sampling_note.strip())

        # ── who coded ──
        doc.add_heading("Coders", level=1)
        doc.add_paragraph("Human coder(s): "
                          + (", ".join(facts["human_raters"]) or "none recorded"))
        if facts["machine_raters"]:
            doc.add_paragraph("Automated coder(s): " + ", ".join(facts["machine_raters"]))
        if facts["demo_codings"]:
            doc.add_paragraph(
                f"WARNING: {facts['demo_codings']} coding(s) in this project come from the "
                f"structured-coding DEMO engine, whose scores are a hash of the text and "
                f"carry no meaning. They are not coding and must not be reported as such. "
                f"Remove them on the Alt Survey Flow page before reporting.")
        if facts["adjudicated_codings"]:
            doc.add_paragraph(
                f"{facts['adjudicated_codings']} coding(s) were reconciled by "
                f"adjudication after a disagreement between coders, and are recorded "
                f"under a separate adjudicated coder.")

        # ── reliability ──
        doc.add_heading("Inter-rater reliability", level=1)
        if not tables:
            doc.add_paragraph(
                "No reliability figure is reported: no two coders coded a shared set "
                "of items in this project. A kappa is only meaningful over items both "
                "raters saw.")
        for label, table in tables.items():
            doc.add_heading(label, level=2)
            _write_table(doc, table)
            for caveat in table.attrs.get("caveats", []):
                doc.add_paragraph(caveat, style="List Bullet")

        doc.add_heading("How to read these figures", level=2)
        for caveat in CAVEATS:
            doc.add_paragraph(caveat, style="List Bullet")

        _write_calibrations(doc, project)

        # ── provenance ──
        doc.add_heading("Codebook and provenance", level=1)
        doc.add_paragraph(
            f"Codebook version {facts['codebook_version'] or 'n/a'} "
            f"({facts['codebook_versions']} snapshot(s) retained). Every change to the "
            f"codebook is snapshotted, so the coding can be traced to the codebook in "
            f"force when it was applied.")
        if models:
            _write_models(doc, models)

        doc.add_heading("Data handling", level=1)
        doc.add_paragraph(PRIVACY_NOTE)

        doc.save(str(out_path))
        return out_path
    finally:
        if close_after:
            project.close()



def _write_calibrations(doc, project) -> None:
    """Structured coding (Alt Survey Flow), reported through its calibration
    sample: per-code precision, recall and kappa against a person, and the
    confidence band the sample supports (analysis.calibration)."""
    from analysis.calibration import CAVEATS as CALIBRATION_CAVEATS, calibrations

    found = calibrations(project)
    if not found:
        return
    doc.add_heading("Structured coding: calibration against a human sample", level=1)
    for calibration in found:
        doc.add_heading(f"Run {calibration['run_id']}", level=2)
        if "__demo_" in calibration["run_id"]:
            doc.add_paragraph(
                "WARNING: this run used the DEMO engine, whose scores are a hash of the "
                "text. These figures describe that hash, not coding.")
        band = calibration["band"]
        doc.add_paragraph(
            f"A random sample of {calibration['n_responses']} response(s) was coded by a "
            f"person on every code. At a target precision of {calibration['target']:.0%}, "
            f"the sample supports auto-assigning at scores of "
            f"{band['high'] if band['high'] is not None else '(no threshold reached)'} and "
            f"above, and treating scores below "
            f"{band['low'] if band['low'] is not None else '(no threshold reached)'} as no "
            f"code. {band['note']}".strip())
        table = calibration["table"].copy()
        for column in ("precision", "recall", "f1", "kappa", "precision_at_high",
                       "recall_at_high"):
            table[column] = table[column].round(3)
        table["recommended_high"] = table["recommended_high"].map(
            lambda v: "not reached" if v is None or v != v else f"{v:.2f}")
        _write_table(doc, table)
    for caveat in CALIBRATION_CAVEATS:
        doc.add_paragraph(caveat, style="List Bullet")


def _write_models(doc, models) -> None:
    """The statistical models, with what qualifies each one.

    Deliberately after reliability and before data handling: a reader works down
    from how the coding was produced, to how much the coders agreed, to what was
    then estimated from it. A model presented before its coding's reliability
    would invite reading the estimate as though the variables were measured.
    """
    from core.model_spec import METHODS, describe

    doc.add_heading("Statistical models", level=1)
    doc.add_paragraph(
        "Each model below was fitted to codes and descriptors from this "
        "project. Codes are variables an analyst constructed, not quantities "
        "measured, and the caveats under each model say what follows from that.")

    for index, entry in enumerate(models, start=1):
        result, request = entry
        method = METHODS.get(getattr(result, "method", ""), None)
        doc.add_heading(
            f"Model {index}: {method.label if method else result.method}",
            level=2)
        doc.add_paragraph(describe(request))

        effect = getattr(result, "effect", None)
        if effect is not None:
            low, high = effect.interval
            paragraph = doc.add_paragraph()
            paragraph.add_run(f"Effect of {effect.treatment}: ").bold = True
            paragraph.add_run(
                f"{effect.coefficient:.3f}, 95% interval "
                f"[{low:.3f}, {high:.3f}], standard error "
                f"{effect.std_error:.3f}.")

        frame = result.frame()
        if not frame.empty:
            _write_table(doc, frame)

        doc.add_paragraph(
            f"Fitted on {result.n_cases} case(s) with seed {result.seed}"
            + (f"; {result.score_name} {result.score:.3f}."
               if result.score is not None else "."))

        for note in result.notes:
            doc.add_paragraph(note, style="List Bullet")

def _write_table(doc, frame: pd.DataFrame) -> None:
    """Render a DataFrame as a Word table with a header row."""
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Light Grid Accent 1"
    for i, column in enumerate(frame.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(column)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for i, column in enumerate(frame.columns):
            value = row[column]
            if isinstance(value, float):
                cells[i].text = "" if pd.isna(value) else f"{value:.3f}"
            else:
                cells[i].text = str(value)
