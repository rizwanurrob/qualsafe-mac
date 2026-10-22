#!/usr/bin/env python
"""
Model exports — Word, Excel, and the design matrix (`V6_PLAN.md` Phase D).

**This is the half that justifies modelling here rather than in SPSS.** Fitting
a lasso beside the coding is a convenience; a coefficient table that arrives in
a chapter still carrying the coding provenance, the seed, the effective n and
the caveats is the argument. Without this the honest recommendation would be to
export a CSV and use Stata.

Three shapes, for three different readers.

**Word** is the publishable table: terms down, the estimate beside them, with
everything qualifying it printed *underneath* rather than left in the software.
A coefficient table looks authoritative whatever produced it, and the person
reading it in a chapter will not have the program open to be warned by it.

**Excel** is the working shape — the estimates, the model specification, and
the caveats on separate sheets — and like every other workbook here it carries
no verbatim response text, so it stays email-safe.

**The design matrix as CSV** is the interoperability escape hatch. Refusing to
let an analyst take their data to R or Stata would be the same mistake the
REFI-QDA work exists to avoid, so the matrix leaves with a sidecar recording
every encoding decision: what a case was, whose codings built the columns, which
level each dummy is relative to, and what was dropped for having no variation.
A matrix without that sidecar is a table of numbers nobody can reconstruct.

What these exports refuse to do
-------------------------------
They do not print a *p*-value or an interval for anything except double lasso,
because nothing else here produces one and inventing a column would be worse
than omitting it. They do not describe a result as causal. And they do not drop
the caveats to make a table fit: if the qualifications will not fit on the page,
the table is the thing that gets shorter.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd

from core.manifest import PACKAGE_VERSION
from core.model_spec import METHODS, ModelRequest, describe

#: Sheet names in the workbook. Short, because Excel truncates at 31 characters
#: and a truncated sheet name is a support question.
SHEET_ESTIMATES = "Estimates"
SHEET_SPEC = "Specification"
SHEET_CAVEATS = "How to read"



class RefusedModelError(ValueError):
    """Raised when an export is asked for a model that was not fitted.

    A refused result carries no estimates and no caveats, so exporting it
    produces a document that *looks* like a model report and contains nothing —
    a heading, a seed, and blank space where the numbers should be. That is
    worse than no file at all, because a file gets circulated. The refusal is
    the finding; it belongs on screen, not in a Word document with a title.
    """


def _require_fitted(result) -> None:
    if not getattr(result, "usable", False):
        raise RefusedModelError(
            "This model was not fitted, so there is nothing to export. "
            + (result.refusal or "No reason was recorded."))


def build_model_docx(
    result,
    request: ModelRequest,
    out_path: Union[str, Path],
    *,
    matrix=None,
    title: str = "Model results",
) -> Path:
    """Write one fitted model as a Word document. Returns the path.

    Raises `RefusedModelError` when the model was refused: see that class.
    """
    _require_fitted(result)

    from docx import Document
    from docx.shared import Pt

    from core.project import SCHEMA_VERSION

    out_path = Path(out_path)
    doc = Document()
    doc.add_heading(title, level=0)

    method = METHODS.get(result.method)
    stamp = doc.add_paragraph()
    run = stamp.add_run(
        f"Generated {datetime.now().strftime('%B %d, %Y')} · QualSafe "
        f"{PACKAGE_VERSION} · project format v{SCHEMA_VERSION} · "
        f"{method.label if method else result.method} · "
        f"{result.n_cases} case(s) · seed {result.seed}")
    run.font.size = Pt(9)

    doc.add_paragraph(describe(request))

    # -- the effect, when there is one --------------------------------------
    if result.effect is not None:
        _write_effect(doc, result.effect)

    # -- the estimates -------------------------------------------------------
    doc.add_heading(
        "Coefficients" if result.kind == "coefficient" else "Variable importance",
        level=1)
    frame = result.frame()
    if frame.empty:
        doc.add_paragraph("The model produced no estimates.")
    else:
        _write_table(doc, frame)

    if result.kind == "importance":
        note = doc.add_paragraph()
        note_run = note.add_run(
            "Importances have no direction and no units: each is how much worse "
            "the model predicted when that column was shuffled. An importance "
            "cannot say whether a code makes the outcome more or less likely.")
        note_run.italic = True

    if result.rules:
        doc.add_heading("Rules", level=1)
        rules = doc.add_paragraph(result.rules)
        rules.style = doc.styles["No Spacing"]
        for paragraph_run in rules.runs:
            paragraph_run.font.name = "Consolas"
            paragraph_run.font.size = Pt(8)

    # -- the fit -------------------------------------------------------------
    doc.add_heading("How the model was fitted", level=1)
    for line in _specification_lines(result, request, matrix):
        doc.add_paragraph(line, style="List Bullet")

    # -- the caveats, never optional ----------------------------------------
    doc.add_heading("How to read this", level=1)
    for caveat in result.notes:
        doc.add_paragraph(caveat, style="List Bullet")
    doc.add_paragraph(
        f"Seed {result.seed}. The same project and the same seed give the same "
        "numbers.", style="List Bullet")

    doc.save(str(out_path))
    return out_path


def _write_effect(doc, effect) -> None:
    """The estimate double lasso exists to produce, above everything else."""
    from docx.shared import Pt

    doc.add_heading(f"Effect of {effect.treatment}", level=1)
    low, high = effect.interval
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{effect.coefficient:.3f}").bold = True
    paragraph.add_run(
        f"   95% interval [{low:.3f}, {high:.3f}]   "
        f"standard error {effect.std_error:.3f}")

    if effect.only_from_treatment:
        listed = ", ".join(effect.only_from_treatment[:6])
        note = doc.add_paragraph()
        run = note.add_run(
            f"Controls kept only because they predict the treatment: {listed}. "
            "A single-selection lasso would have dropped these, and dropping "
            "them is what biases the estimate.")
        run.italic = True
        run.font.size = Pt(9)


def _specification_lines(result, request: ModelRequest, matrix) -> List[str]:
    """Everything a reader needs to reconstruct the fit.

    This is the list that makes the export worth having. A reviewer holding the
    document should be able to rebuild the model without the project open.
    """
    method = METHODS.get(result.method)
    lines = [
        f"Method: {method.label if method else result.method}.",
        f"Outcome: {result.outcome}.",
    ]
    if request.treatment.first:
        lines.append(f"Treatment: {request.treatment.first}.")
    lines.append(
        f"Terms on the right-hand side: {len(request.right_hand_side())} "
        f"({', '.join(request.right_hand_side()) or 'none'}).")
    if request.controls.fields:
        lines.append("Forced controls, never penalised out: "
                     + ", ".join(request.controls.fields) + ".")
    lines.append(f"Cases modelled: {result.n_cases}"
                 + (f", from {result.n_cases_before} before dropping cases with "
                    "missing values." if result.n_cases_before > result.n_cases
                    else "."))
    if result.score is not None:
        lines.append(f"{result.score_name}: {result.score:.3f}, "
                     f"over {request.folds} folds.")
    for key, value in sorted(result.chosen.items()):
        lines.append(f"Chosen by the fit — {key.replace('_', ' ')}: {value}.")

    if matrix is not None:
        lines.append(f"Design matrix: {matrix.describe()}.")
        if matrix.reference_levels:
            pairs = "; ".join(f"{key} against “{level}”"
                              for key, level in sorted(matrix.reference_levels.items()))
            lines.append(f"Categorical reference levels: {pairs}.")
        if matrix.dropped:
            lines.append(
                f"Columns dropped before fitting for too little variation: "
                + ", ".join(sorted(matrix.dropped)) + ".")
    lines.append(f"Random seed: {result.seed}.")
    return lines


def _write_table(doc, frame: pd.DataFrame) -> None:
    """Render a DataFrame as a Word table with a header row."""
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Light Grid Accent 1"
    for index, column in enumerate(frame.columns):
        cell = table.rows[0].cells[index]
        cell.text = str(column)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(frame.columns):
            value = row[column]
            if isinstance(value, float):
                cells[index].text = "" if pd.isna(value) else f"{value:.4f}"
            else:
                cells[index].text = str(value)


def build_model_workbook(
    result,
    request: ModelRequest,
    out_path: Union[str, Path],
    *,
    matrix=None,
) -> Path:
    """Write the model to an Excel workbook. Returns the path.

    Three sheets, because the estimates, the specification and the caveats are
    read by different people at different times and merging them into one sheet
    means the caveats get scrolled past.

    Raises `RefusedModelError` when the model was refused.
    """
    _require_fitted(result)

    out_path = Path(out_path)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        frame = result.frame()
        if frame.empty:
            frame = pd.DataFrame({"term": [], result.kind: []})
        frame.to_excel(writer, sheet_name=SHEET_ESTIMATES, index=False)

        specification = pd.DataFrame(
            {"detail": _specification_lines(result, request, matrix)})
        specification.to_excel(writer, sheet_name=SHEET_SPEC, index=False)

        notes = list(result.notes) + [
            f"Seed {result.seed}. The same project and the same seed give the "
            "same numbers."]
        pd.DataFrame({"caveat": notes}).to_excel(
            writer, sheet_name=SHEET_CAVEATS, index=False)
    return out_path


def export_design_matrix(
    matrix,
    out_path: Union[str, Path],
    *,
    include_sidecar: bool = True,
) -> Path:
    """Write the design matrix as CSV, with a sidecar describing it.

    The interoperability escape hatch. Someone who wants R, Stata or SPSS should
    be able to have the data — refusing would be the same mistake the REFI-QDA
    work exists to avoid — but a matrix of `code::Cost` columns with no
    explanation is a table nobody can reconstruct. The sidecar records what a
    case was, whose codings built the columns, which level each dummy is
    relative to, and what was dropped.
    """
    out_path = Path(out_path)
    matrix.frame.to_csv(out_path, index=False)
    if include_sidecar:
        sidecar = out_path.with_suffix(".json")
        sidecar.write_text(json.dumps(design_matrix_sidecar(matrix), indent=2),
                           encoding="utf-8")
    return out_path


def design_matrix_sidecar(matrix) -> Dict[str, Any]:
    """Everything needed to read the CSV correctly, as plain JSON."""
    from core.project import SCHEMA_VERSION

    return {
        "produced_by": f"QualSafe {PACKAGE_VERSION}",
        "project_format": SCHEMA_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "unit": matrix.unit,
        "case_count": matrix.n_cases,
        "code_encoding": matrix.encoding,
        "rater": matrix.rater_id,
        "raters_pooled": list(matrix.raters_pooled),
        "reference_levels": dict(matrix.reference_levels),
        "numeric_descriptors": dict(matrix.numeric_descriptors),
        "dropped_columns": dict(matrix.dropped),
        "code_columns": matrix.code_columns(),
        "descriptor_columns": matrix.descriptor_columns(),
        "caveats": matrix.caveats(),
        "notes": [
            "A code column is 0/1 or a count and is never missing: a code that "
            "was not applied to a case is a real zero.",
            "A descriptor column CAN be missing, and missing means not "
            "collected. Do not fill it with zero.",
            "Categorical descriptors are one-hot encoded with one level held "
            "out; see reference_levels for which.",
        ],
    }
