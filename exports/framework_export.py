#!/usr/bin/env python
"""
Framework matrix exports — Word and Excel (v4, Phase 5.3).

The framework matrix is a synthesis *deliverable*: the grid and its cell
summaries are what goes into a report, not something to read only on screen.

Two shapes, because they serve different readers. **Word** renders the matrix as
a narrative — one section per case, its themes and summaries beneath — which is
how a findings chapter is actually written. **Excel** renders it as a grid, which
is how a team works through coverage and spots the empty cells.

Cell summaries are the analyst's own writing (or an LLM draft they accepted), so
they are exported as written. Supporting *passages* are verbatim response text
and are therefore opt-in, consistent with every other export here.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple, Union

import pandas as pd

from analysis.framework import (
    cell_excerpt_texts, framework_cells, framework_grid, get_cell_summary,
)
from core.manifest import PACKAGE_VERSION


def _case_labels(project, include_uncased: bool) -> Dict[str, str]:
    from analysis.framework import UNCASED

    labels = {c.case_id: c.name for c in project.list_cases()}
    if include_uncased:
        labels[UNCASED] = UNCASED
    return labels


def collect_matrix(
    project, *, code_names: Sequence[str], include_uncased: bool = False,
) -> Tuple[pd.DataFrame, Dict[Tuple[str, str], List[str]], Dict[str, str]]:
    """The grid, the per-cell excerpt ids, and the case labels."""
    coding = project.coding_spans_table()
    cells = framework_cells(coding, code_names=code_names,
                            include_uncased=include_uncased)
    labels = _case_labels(project, include_uncased)
    grid = framework_grid(cells, labels, list(code_names))
    return grid, cells, labels


def build_framework_docx(
    project,
    out_path: Union[str, Path],
    *,
    code_names: Sequence[str],
    include_uncased: bool = False,
    include_passages: bool = False,
    passages_per_cell: int = 3,
    title: str = "Framework Matrix",
) -> Path:
    """Write the framework matrix as a Word document. Returns the path."""
    from docx import Document
    from docx.shared import Pt

    from core.project import Project

    close_after = False
    if not hasattr(project, "coding_spans_table"):
        project = Project.open(project)
        close_after = True

    out_path = Path(out_path)
    try:
        grid, cells, labels = collect_matrix(
            project, code_names=code_names, include_uncased=include_uncased)

        doc = Document()
        doc.add_heading(title, level=0)
        stamp = doc.add_paragraph()
        run = stamp.add_run(
            f"Generated {datetime.now().strftime('%B %d, %Y')} · QualSafe "
            f"{PACKAGE_VERSION} · {len(grid)} case(s) × {len(code_names)} theme(s) · "
            f"{'with' if include_passages else 'without'} supporting quotes")
        run.font.size = Pt(9)

        # Coverage grid first — what is and isn't covered, at a glance.
        doc.add_heading("Coverage", level=1)
        table = doc.add_table(rows=1, cols=len(grid.columns))
        table.style = "Light Grid Accent 1"
        for i, column in enumerate(grid.columns):
            cell = table.rows[0].cells[i]
            cell.text = str(column)
            for paragraph in cell.paragraphs:
                for r in paragraph.runs:
                    r.bold = True
        for _, row in grid.iterrows():
            cells_row = table.add_row().cells
            for i, column in enumerate(grid.columns):
                cells_row[i].text = str(row[column])

        # Then the narrative: one section per case.
        for case_id, label in sorted(labels.items(), key=lambda kv: kv[1].lower()):
            has_any = any((case_id, code) in cells for code in code_names)
            if not has_any:
                continue
            doc.add_heading(label, level=1)
            for code in code_names:
                excerpt_ids = cells.get((case_id, code), [])
                if not excerpt_ids:
                    continue
                doc.add_heading(code, level=2)
                memo = get_cell_summary(project, case_id, code)
                if memo is not None and memo.text.strip():
                    doc.add_paragraph(memo.text.strip())
                else:
                    empty = doc.add_paragraph(
                        f"(no summary written; {len(excerpt_ids)} passage(s) coded)")
                    empty.runs[0].italic = True
                if include_passages:
                    for text in cell_excerpt_texts(project, excerpt_ids)[:passages_per_cell]:
                        doc.add_paragraph(f"“{text.strip()[:400]}”", style="Intense Quote")

        doc.save(str(out_path))
        return out_path
    finally:
        if close_after:
            project.close()


def build_framework_workbook(
    project,
    out_path: Union[str, Path],
    *,
    code_names: Sequence[str],
    include_uncased: bool = False,
) -> Path:
    """Write the matrix to Excel: a counts grid and a summaries sheet.

    Never contains verbatim response text — like `exports/excel_workbook`, this
    is the email-safe shape.
    """
    from core.project import Project

    close_after = False
    if not hasattr(project, "coding_spans_table"):
        project = Project.open(project)
        close_after = True

    out_path = Path(out_path)
    try:
        grid, cells, labels = collect_matrix(
            project, code_names=code_names, include_uncased=include_uncased)

        summaries = []
        for (case_id, code), excerpt_ids in sorted(cells.items()):
            memo = get_cell_summary(project, case_id, code)
            summaries.append({
                "case": labels.get(case_id, case_id),
                "theme": code,
                "passages": len(excerpt_ids),
                "summary": memo.text if memo is not None else "",
            })
        summary_frame = pd.DataFrame(
            summaries, columns=["case", "theme", "passages", "summary"])

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            grid.to_excel(writer, sheet_name="Coverage", index=False)
            if summary_frame.empty:
                pd.DataFrame({"note": ["No cells with coded passages."]}).to_excel(
                    writer, sheet_name="Summaries", index=False)
            else:
                summary_frame.to_excel(writer, sheet_name="Summaries", index=False)
        return out_path
    finally:
        if close_after:
            project.close()
