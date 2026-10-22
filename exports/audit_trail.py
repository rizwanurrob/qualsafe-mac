#!/usr/bin/env python
"""
The audit trail: every dated analytic decision in a project, in order (U3).

Lincoln and Guba's *confirmability audit* asks whether a stranger can follow
the trail from a finding back to the data and the decisions made on the way:
how the codebook changed and why, who coded what and when, what was
reconciled, what was removed. Reviewers and ethics boards ask for it by name.
QualSafe stores every ingredient; this module lays them out as one
chronological record, as a Word document for reading and an Excel workbook
for filtering.

Sources, and what each can and cannot say:

* **The event log** (store format v9+): codebook edits with before and after,
  memos, adjudication, imports, merges, auto-coding, erasure. Complete from
  the moment a project was first opened in 4.13; nothing earlier.
* **Records that were always dated**, which fill the earlier period: analysis
  runs, codebook snapshots, review records (v8+, summarised per coder per
  day), structured-coding reviews, documents marked finished, backups, and
  the erasure log file beside the store.

Redaction: **never response text, never response ids.** Memo text only when
the analyst opts in, because a memo can quote a response. Document titles are
not shown (a file name can be identifying); documents are counted.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd

from core.manifest import PACKAGE_VERSION

COLUMNS = ["at", "who", "what", "detail", "source"]

ABOUT = [
    "Every row is a dated decision or event in this project, oldest first.",
    "From the first time the project was opened in QualSafe 4.13, decisions are "
    "recorded as they are made ('event log'). Before that, this record is "
    "reconstructed from what was always dated: analysis runs, codebook snapshots, "
    "coding reviews, backups and erasures. Edits made before 4.13 that left no "
    "dated record (a code renamed, a definition changed, a memo written) are not "
    "shown, because they cannot be dated.",
    "No response text and no response identifiers appear here. Erasures are "
    "recorded by a one-way hash. Memo text appears only if the person who built "
    "this file chose to include it.",
]


def _row(at: str, who: str, what: str, detail: str, source: str) -> Dict[str, str]:
    return {"at": at or "", "who": who or "", "what": what, "detail": detail, "source": source}


def _detail_text(detail: Dict[str, Any]) -> str:
    parts = []
    for key, value in detail.items():
        if value in (None, "", [], {}):
            continue
        if isinstance(value, (list, dict)):
            value = json.dumps(value, ensure_ascii=False)
        parts.append(f"{key.replace('_', ' ')}: {value}")
    return "; ".join(parts)


def audit_trail_frame(project, *, project_dir: Optional[Union[str, Path]] = None,
                      include_memo_text: bool = False) -> pd.DataFrame:
    """Every dated decision in the project, oldest first. See the module doc."""
    rows: List[Dict[str, str]] = []

    # 1. the event log
    for event in project.list_events():
        rows.append(_row(event["at"], event["actor"], event["summary"],
                         _detail_text(event["detail"]), "event log"))

    # 2. analysis runs
    for run in project.list_runs():
        thresholds = {k: v for k, v in (run.thresholds or {}).items() if v is not None}
        rows.append(_row(run.timestamp, "", f"Analysis run: {run.engine or 'engine'} "
                                            f"({run.model or 'model n/a'})",
                         _detail_text({"mode": run.mode, "QualSafe": run.package_version,
                                       "thresholds": thresholds, "counts": run.counts}),
                         "runs"))

    # 3. codebook snapshots (the before-state of each structural change)
    for version in project.list_codebook_versions():
        trigger = version.get("trigger") or "snapshot"
        rows.append(_row(version["created_at"], "",
                         f"Codebook snapshot ({trigger})",
                         _detail_text({"note": version.get("note"),
                                       "codebook hash": version.get("codebook_hash")}),
                         "codebook versions"))

    conn = project._conn
    # 4. human review sessions: one row per coder per day (v8 review record)
    human = {r.rater_id for r in project.list_raters() if r.kind == "human"}
    for r in conn.execute(
            "SELECT rater_id, substr(reviewed_at, 1, 10) AS day, MIN(reviewed_at) AS first, "
            "COUNT(*) AS n FROM rater_coverage WHERE reviewed_at != '' "
            "GROUP BY rater_id, day ORDER BY first"):
        if r["rater_id"] in human:
            rows.append(_row(r["first"], r["rater_id"],
                             f"Coding recorded: {r['n']} response(s) reviewed",
                             "including any left with no code", "review record"))

    # 5. structured-coding reviews (Alt Survey Flow)
    for r in conn.execute(
            "SELECT reviewer, substr(reviewed_at, 1, 10) AS day, MIN(reviewed_at) AS first, "
            "SUM(review_status = 'confirmed') AS confirmed, SUM(review_status = 'rejected') AS rejected, "
            "SUM(review_status = 'added_by_human') AS added, SUM(review_status = 'flagged') AS flagged "
            "FROM coding_decisions WHERE reviewed_at != '' GROUP BY reviewer, day ORDER BY first"):
        rows.append(_row(r["first"], r["reviewer"], "Structured-coding suggestions reviewed",
                         _detail_text({"confirmed": r["confirmed"], "rejected": r["rejected"],
                                       "added": r["added"], "flagged": r["flagged"]}),
                         "structured coding"))

    # 6. documents marked finished (counted, not named: a title can identify)
    for r in conn.execute(
            "SELECT rater_id, substr(marked_at, 1, 10) AS day, MIN(marked_at) AS first, COUNT(*) AS n "
            "FROM document_coverage GROUP BY rater_id, day ORDER BY first"):
        rows.append(_row(r["first"], r["rater_id"],
                         f"{r['n']} document(s) marked as fully coded", "", "review record"))

    # 7. memos: dated ones are in the event log; say how many predate dating
    memos = project.list_memos()
    undated = [m for m in memos if not m.created_at]
    if undated:
        rows.append(_row("", "", f"{len(undated)} memo(s) written before memos were dated",
                         "", "memos"))
    if include_memo_text:
        for memo in memos:
            rows.append(_row(memo.updated_at or memo.created_at, memo.author,
                             f"Memo ({memo.target_type})", memo.text, "memo text"))

    # 8. files beside the store: the erasure log and the backups
    base = Path(project_dir) if project_dir else Path(project.path).parent
    logged_hashes = {e["target_id"] for e in project.list_events()
                     if e["kind"] == "respondent_erased"}
    log = base / "audit_log.jsonl"
    if log.exists():
        for line in log.read_text(encoding="utf-8").splitlines():
            try:
                entry = json.loads(line)
            except ValueError:
                continue
            if entry.get("erased_response_id_hash") in logged_hashes:
                continue      # already in the event log
            rows.append(_row(entry.get("timestamp", ""), "", "A respondent's data was erased",
                             _detail_text({"response id hash": entry.get("erased_response_id_hash"),
                                           "backups purged": entry.get("backups_purged")}),
                             "erasure log"))
    from core.backups import list_backups
    for backup in list_backups(project.path):
        match = re.match(r"project_(\d{8})_(\d{6})_\d+_\d+_(.+)\.qnw$", backup.name)
        if match:
            day, clock, reason = match.groups()
            at = (f"{day[:4]}-{day[4:6]}-{day[6:]}T"
                  f"{clock[:2]}:{clock[2:4]}:{clock[4:]}")
            rows.append(_row(at, "", f"Backup taken before {reason.replace('_', ' ')}",
                             "", "backups"))

    frame = pd.DataFrame(rows, columns=COLUMNS)
    # Undated rows first (they predate everything dated), then chronological.
    frame["_undated"] = frame["at"] == ""
    frame = frame.sort_values(["_undated", "at"], ascending=[False, True], kind="stable")
    return frame.drop(columns="_undated").reset_index(drop=True)


def _open(project):
    from core.project import Project

    if hasattr(project, "list_events"):
        return project, False
    return Project.open(project), True


def build_audit_trail_xlsx(project, out_path: Union[str, Path], *,
                           include_memo_text: bool = False) -> Path:
    """Excel: the trail (filterable), and a sheet saying what it is and is not."""
    from analyze_survey import sanitize_frame_for_csv

    project, close_after = _open(project)
    out_path = Path(out_path)
    try:
        frame = audit_trail_frame(project, include_memo_text=include_memo_text)
        with pd.ExcelWriter(out_path) as writer:
            sanitize_frame_for_csv(frame).to_excel(writer, sheet_name="Audit trail", index=False)
            pd.DataFrame({"about this record": ABOUT + [
                f"Generated {datetime.now().isoformat(timespec='seconds')} by QualSafe "
                f"{PACKAGE_VERSION}."]}).to_excel(writer, sheet_name="About", index=False)
    finally:
        if close_after:
            project.close()
    return out_path


def build_audit_trail_docx(project, out_path: Union[str, Path], *,
                           include_memo_text: bool = False,
                           title: str = "Audit trail") -> Path:
    """Word: what the record is, then the trail as a dated table."""
    from docx import Document
    from docx.shared import Pt

    from core.project import SCHEMA_VERSION

    project, close_after = _open(project)
    out_path = Path(out_path)
    try:
        frame = audit_trail_frame(project, include_memo_text=include_memo_text)
        doc = Document()
        doc.add_heading(title, level=0)
        stamp = doc.add_paragraph().add_run(
            f"Generated {datetime.now().strftime('%B %d, %Y')} · QualSafe {PACKAGE_VERSION} · "
            f"project format v{SCHEMA_VERSION}")
        stamp.font.size = Pt(9)
        doc.add_heading("What this record is", level=1)
        for line in ABOUT:
            doc.add_paragraph(line, style="List Bullet")
        doc.add_heading(f"The record ({len(frame)} entries)", level=1)
        if frame.empty:
            doc.add_paragraph("Nothing dated has happened in this project yet.")
        else:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            for cell, label in zip(table.rows[0].cells, ("When", "Who", "What", "Detail")):
                cell.text = label
            for _, r in frame.iterrows():
                cells = table.add_row().cells
                cells[0].text = r["at"].replace("T", " ") or "undated"
                cells[1].text = r["who"]
                cells[2].text = r["what"]
                cells[3].text = r["detail"]
        doc.save(out_path)
    finally:
        if close_after:
            project.close()
    return out_path
